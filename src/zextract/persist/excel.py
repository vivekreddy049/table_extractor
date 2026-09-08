"""S10 -- one .xlsx per logical table, four sheets.

The typing policy is the point of this module. Native Excel numbers are written
ONLY where confidence clears the configured band; below it the raw string is
written and the cell is tinted. A spreadsheet where every cell looks equally
authoritative is the silent-error failure transplanted into Excel -- the reader
has no way to tell which figures to check. A file with mixed types in one column
is the honest representation of a document we are not uniformly sure about, and
the Issues sheet names every instance.
"""

from __future__ import annotations

import datetime as _dt
import io
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment as XlAlignment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import Config
from ..model import Document, Page, Table, ValueType
from ..s7_normalise.values import NUMERIC_TYPES

_HEADER_FILL = PatternFill("solid", fgColor="EFEFEA")
_LOW_FILL = PatternFill("solid", fgColor="FDF3E3")
_ERR_FILL = PatternFill("solid", fgColor="FAE7E1")
_BOLD = Font(bold=True)
_SLUG = re.compile(r"[^a-z0-9]+")


def slug(text: str, limit: int = 40) -> str:
    s = _SLUG.sub("-", text.lower()).strip("-")
    return (s[:limit].rstrip("-") or "table")


def table_filename(table: Table, logical_index: int) -> str:
    label = ""
    if table.rows:
        first = table.cell_at(0, 0)
        if first and first.raw_text.strip():
            label = first.raw_text
    return (
        f"T{logical_index:03d}__p{table.start_page}-{table.end_page}__{slug(label)}.xlsx"
    )


def write_workbook(
    doc: Document, page: Page, table: Table, logical_index: int, cfg: Config, out_dir: Path
) -> Path:
    native_min = cfg.f("excel.native_type_min_confidence")

    wb = Workbook()
    _sheet_data(wb.active, table, native_min)
    _sheet_context(wb.create_sheet("Context"), doc, page, table, logical_index, cfg)
    _sheet_provenance(wb.create_sheet("Provenance"), table)
    _sheet_issues(wb.create_sheet("Issues"), table, native_min)

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / table_filename(table, logical_index)
    _save_deterministic(wb, path)
    return path


# A fixed instant, not "now". The brief diffs two runs of the Excel output
# byte-for-byte, and openpyxl otherwise stamps the current time into
# docProps/core.xml and into every zip entry header.
_EPOCH = _dt.datetime(2020, 1, 1, 0, 0, 0)
_EPOCH_TEXT = "2020-01-01T00:00:00Z"
_ZIP_DATE = (1980, 1, 1, 0, 0, 0)
_TIMESTAMP_RE = re.compile(
    r"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)"
)


def _save_deterministic(wb: Workbook, path: Path) -> None:
    """Write an .xlsx whose bytes depend only on its content.

    Two things leak wall-clock time into a workbook and both have to go:
    the document properties, and the per-entry timestamps in the zip container.
    The file is therefore built in memory and then repacked with fixed dates and
    a fixed entry order.
    """
    wb.properties.creator = "zextract"
    wb.properties.lastModifiedBy = "zextract"
    wb.properties.created = _EPOCH
    wb.properties.modified = _EPOCH

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    with zipfile.ZipFile(buf) as src:
        entries = sorted(src.namelist())
        payload = {name: src.read(name) for name in entries}

    # openpyxl overwrites properties.modified with the current UTC time inside
    # save(), ignoring whatever was set beforehand, so the timestamp has to be
    # rewritten after the fact rather than configured before it. This was the
    # last non-determinism in the Excel output.
    core = payload.get("docProps/core.xml")
    if core is not None:
        payload["docProps/core.xml"] = _TIMESTAMP_RE.sub(
            lambda m: f"{m.group(1)}{_EPOCH_TEXT}{m.group(3)}", core.decode("utf-8")
        ).encode("utf-8")

    with open(path, "wb") as fh:
        with zipfile.ZipFile(fh, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as dst:
            for name in entries:
                info = zipfile.ZipInfo(name, date_time=_ZIP_DATE)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                dst.writestr(info, payload[name])


def _sheet_data(ws, table: Table, native_min: float) -> None:
    ws.title = "Data"
    for c in table.cells:
        cell = ws.cell(row=c.row_idx + 1, column=c.col_idx + 1)
        confident = c.confidence >= native_min

        if c.value_type is ValueType.ERROR:
            # Never coerced to 0 or NULL: the source really does say #REF!.
            cell.value = c.raw_text
            cell.fill = _ERR_FILL
        elif c.value_type in NUMERIC_TYPES and c.normalized_value is not None and confident:
            cell.value = c.normalized_value
            if c.value_type is ValueType.PERCENT:
                cell.number_format = "0.00%"
            elif c.value_type is ValueType.CURRENCY:
                cell.number_format = "#,##0.00"
        elif c.value_type is ValueType.DATE and c.normalized_text and confident:
            cell.value = c.normalized_text
        else:
            cell.value = c.raw_text
            if c.raw_text.strip() and not confident:
                cell.fill = _LOW_FILL

        cell.alignment = XlAlignment(
            horizontal="right" if c.value_type in NUMERIC_TYPES else "left",
            wrap_text=True,
            vertical="top",
        )

    for c in table.cells:
        if (c.row_span > 1 or c.col_span > 1) and c.raw_text.strip():
            try:
                ws.merge_cells(
                    start_row=c.row_idx + 1,
                    start_column=c.col_idx + 1,
                    end_row=c.row_idx + c.row_span,
                    end_column=c.col_idx + c.col_span,
                )
            except ValueError:
                continue

    if table.rows:
        for col in table.columns:
            ws.cell(row=1, column=col.index + 1).font = _BOLD
            ws.cell(row=1, column=col.index + 1).fill = _HEADER_FILL
        ws.freeze_panes = "A2"

    for col in table.columns:
        ws.column_dimensions[get_column_letter(col.index + 1)].width = 26


def _sheet_context(ws, doc: Document, page: Page, table: Table, logical_index: int, cfg: Config) -> None:
    rows = [
        ("document", doc.filename),
        ("sha256", doc.sha256),
        ("logical_index", logical_index),
        ("page range", f"{table.start_page}-{table.end_page}"),
        ("rows x cols", f"{len(table.rows)} x {len(table.columns)}"),
        ("partition source", table.partition_source),
        ("partition support", round(table.support, 4)),
        ("flags", ", ".join(sorted(table.flags)) or "none"),
        ("title", table.title or "NOT EXTRACTED"),
        ("caption", table.caption or "NOT EXTRACTED"),
        ("section path", table.section_path or "NOT EXTRACTED"),
        ("unit / scale note", table.unit_scale_note or "NOT EXTRACTED"),
        (
            "unit / scale source",
            table.unit_scale_source or "NOT EXTRACTED",
        ),
        (
            "footnotes",
            "; ".join(
                f"{fn.marker}: {fn.text}"
                for fn in doc.footnotes
                if any(fn.index in c.footnote_ids for c in table.cells)
            )
            or "none",
        ),
        (
            "linked assets",
            ", ".join(
                f"{link.relation} asset#{link.asset_index} ({link.confidence:.2f})"
                for link in table.asset_links
            )
            or "none",
        ),
        (
            "table confidence",
            round(sum(c.confidence for c in table.cells) / len(table.cells), 4)
            if table.cells
            else "n/a",
        ),
        (
            "confidence basis",
            "calibrated (frozen model_v1.json); see LIMITATIONS.md §2"
            if cfg.b("confidence.apply_calibration")
            else "UNCALIBRATED - ranking only, see LIMITATIONS.md",
        ),
    ]
    for c in table.columns:
        rows.append(
            (
                f"column {c.index}",
                f"type={c.inferred_type.value} unit={c.unit or '-'} "
                f"align={c.alignment.value} agreement={c.type_agreement:.2f}",
            )
        )
    _write_rows(ws, ("field", "value"), rows)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 72


def _sheet_provenance(ws, table: Table) -> None:
    header = (
        "row_idx", "col_idx", "raw_text", "normalized_value", "normalized_text",
        "value_type", "unit", "scale", "parse_rule", "page_no",
        "bbox_x0", "bbox_y0", "bbox_x1", "bbox_y1", "source", "confidence", "codes",
    )
    rows = []
    for c in sorted(table.cells, key=lambda c: (c.row_idx, c.col_idx)):
        b = c.bbox
        rows.append(
            (
                c.row_idx, c.col_idx, c.raw_text, c.normalized_value, c.normalized_text,
                c.value_type.value, c.unit, c.scale, c.parse_rule, table.page_no,
                b.x0 if b else None, b.y0 if b else None,
                b.x1 if b else None, b.y1 if b else None,
                c.source.value, round(c.confidence, 4), ",".join(c.codes),
            )
        )
    _write_rows(ws, header, rows)


def _sheet_issues(ws, table: Table, native_min: float) -> None:
    header = ("severity", "code", "row_idx", "col_idx", "raw_text", "confidence", "message")
    rows = []
    for i in table.issues:
        cell = (
            table.cell_at(i.row_idx, i.col_idx)
            if i.row_idx is not None and i.col_idx is not None
            else None
        )
        rows.append(
            (
                i.severity, i.code, i.row_idx, i.col_idx,
                cell.raw_text if cell else "",
                round(cell.confidence, 4) if cell else "",
                i.message,
            )
        )
    for c in sorted(table.cells, key=lambda c: (c.row_idx, c.col_idx)):
        if c.confidence < native_min and c.raw_text.strip():
            rows.append(
                (
                    "info", "LOW_CONFIDENCE", c.row_idx, c.col_idx, c.raw_text,
                    round(c.confidence, 4),
                    "written as raw text, not a native Excel value",
                )
            )
    _write_rows(ws, header, rows)
    ws.column_dimensions["G"].width = 60


def _write_rows(ws, header, rows) -> None:
    for j, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
    for i, row in enumerate(rows, start=2):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)
    ws.freeze_panes = "A2"
